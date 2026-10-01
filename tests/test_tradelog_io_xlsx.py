"""Tests for the consolidated-ledger ``.xlsx`` loader and its MCP routing.

The workbook is the four-tab trade log (``Index Options`` / ``Other Options`` /
``Stock & ETF`` / ``Other Transactions``) that merges IBKR and E*Trade
trades.  Its conventions differ from every other loader's input, and these
tests pin the ones that carry real money:

* Quantity is always positive in the sheet, but the P&L engine's
  expire-inference needs it signed (Buy +, Sell -).
* Net Amount is already signed and is taken verbatim.
* Commission is negative in the sheet; the standard frame carries the
  magnitude exactly as recorded, and gross is net with the commission added
  back.  IBKR's column includes all fees; E*Trade's is only the broker's
  charge, because exchange fees are already in the recorded fill price.
* E*Trade option symbols are not OCC; they are rebuilt from Description.

Every row below is invented: fictional tickers, round prices, made-up dates
and a placeholder account id.  Never copy rows, account ids, descriptions,
amounts or file names from a real statement or trade log into this file.
"""
from __future__ import annotations

import base64
import io
import subprocess
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from spx_trade_desk.mcp.server import (
    _load_and_merge,
    _load_file_from_path,
    compute_daily_pnl,
    get_transaction_summary,
)
from spx_trade_desk.tradelog.io.load_etrade_csv import load_transactions_etrade_csv
from spx_trade_desk.tradelog.io.load_xlsx import load_transactions_xlsx
from tests.test_tradelog_io_qfx_pdf import SYNTHETIC_QFX

HEADER = [
    "Date", "Source", "Account", "Underlying", "Symbol", "Description",
    "Order Type", "Quantity", "Price", "Commission", "Net Amount",
]

IBKR = ("IBKR", "U***12345")
ETRADE = ("ETRADE", "ET-ACCT-1")

STANDARD_COLUMNS = [
    "activity_date", "account_id", "description", "transaction_type",
    "symbol", "quantity", "price", "gross_amount", "commission",
    "net_amount", "source_row",
]

# Symbols used in more than one place below.
IBKR_SHORT = "SPXW  250303P05000000"
IBKR_LONG = "SPXW  250303P04950000"
ETRADE_SHORT = "SPXW  250310P05100000"
ETRADE_LONG = "SPXW  250310P05050000"
ETRADE_CLOSE = "SPXW  250310P05075000"
ASSIGNED = "SPXW  250324P05000000"
OTHER_OPTION = "ABCD  250328P00100000"

INDEX_ROWS = [
    # IBKR: OCC symbol already, positive qty, signed net, negative commission.
    (date(2025, 3, 3), *IBKR, "SPXW", IBKR_SHORT,
     "SPXW 03MAR25 5000 P", "Sell", 1, 0.50, -1.00, 49.00),
    (date(2025, 3, 3), *IBKR, "SPXW", IBKR_LONG,
     "SPXW 03MAR25 4950 P", "Buy", 1, 0.20, -1.00, -21.00),
    # E*Trade: non-OCC symbol, option detail lives in Description.  The net
    # (78.00) is below qty x price x 100 (80.00) by more than the commission
    # column (1.00): exchange fees are already in the recorded figures.
    (date(2025, 3, 10), *ETRADE, "SPXW", "SPXW MAR 10 '25 $5100 PUT",
     "PUT  SPXW   03/10/25  5100.000", "Sell To Open", 2, 0.40, -1.00, 78.00),
    (date(2025, 3, 10), *ETRADE, "SPXW", "SPXW MAR 10 '25 $5050 PUT",
     "PUT  SPXW   03/10/25  5050.000", "Buy Open", 2, 0.10, -1.00, -21.00),
    (date(2025, 3, 10), *ETRADE, "SPXW", "SPXW MAR 10 '25 $5075 PUT",
     "PUT  SPXW   03/10/25  5075.000", "Buy To Close", 1, 0.30, -0.50, -30.50),
    # Assignment settlement: no quantity / price / commission.
    (date(2025, 3, 24), *IBKR, "SPXW", ASSIGNED,
     "Option Cash Settlement for: Assignment ( SPXW 24MAR25 5000 P )",
     "Cash Settlement", None, None, None, -25.00),
]
OTHER_OPTION_ROWS = [
    (date(2025, 3, 20), *IBKR, "ABCD", OTHER_OPTION,
     "ABCD 28MAR25 100 P", "Sell", 3, 0.50, -1.50, 148.50),
]
STOCK_ROWS = [
    (date(2025, 3, 5), *ETRADE, "ACME", "ACME",
     "ACME CORP UNSOLICITED TRADE", "Buy", 2, 100, 0, -200.00),
    (date(2025, 3, 18), *ETRADE, "BETA", "BETA",
     "BETA INC UNSOLICITED TRADE", "Sell", 20, 50, -5.00, 995.00),
]
OTHER_TX_ROWS = [
    (date(2025, 3, 6), *IBKR, "FUND", "FUND",
     "FUND Cash Dividend USD 0.25 per Share", "Dividend", None, None, None, 100.00),
    (date(2025, 3, 4), *IBKR, None, "-",
     "Monthly market data fee", "Other Fee", None, None, None, -10.00),
]
TOTAL_ROWS = (
    len(INDEX_ROWS) + len(OTHER_OPTION_ROWS) + len(STOCK_ROWS) + len(OTHER_TX_ROWS)
)  # 11


def _sheet(wb: Workbook, title: str, rows: list, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(HEADER)
    for row in rows:
        ws.append(list(row))
    return ws


def _build_workbook() -> Workbook:
    wb = Workbook()
    index_ws = _sheet(wb, "Index Options", INDEX_ROWS, first=True)
    # Dashboard cells beside the ledger (columns L onward).
    index_ws["M1"] = "helper date"
    index_ws["N1"] = "helper total"
    index_ws["M2"] = date(2025, 3, 3)
    index_ws["N2"] = "=SUM(K2:K3)"
    index_ws["Q1"] = "dashboard note"
    _sheet(wb, "Other Options", OTHER_OPTION_ROWS)
    _sheet(wb, "Stock & ETF", STOCK_ROWS)
    _sheet(wb, "Other Transactions", OTHER_TX_ROWS)
    return wb


@pytest.fixture
def xlsx_bytes() -> bytes:
    buf = io.BytesIO()
    _build_workbook().save(buf)
    return buf.getvalue()


@pytest.fixture
def xlsx_path(tmp_path, xlsx_bytes):
    path = tmp_path / "trade_log.xlsx"
    path.write_bytes(xlsx_bytes)
    return path


@pytest.fixture
def df(xlsx_path) -> pd.DataFrame:
    return load_transactions_xlsx(xlsx_path)


def _row(frame: pd.DataFrame, **match) -> pd.Series:
    mask = pd.Series(True, index=frame.index)
    for col, val in match.items():
        mask &= frame[col] == val
    assert mask.sum() == 1, f"expected exactly one row for {match}, got {mask.sum()}"
    return frame[mask].iloc[0]


# ---------------------------------------------------------------------------
# Loader: shape
# ---------------------------------------------------------------------------

class TestLoaderShape:
    def test_columns_are_the_standard_schema(self, df):
        assert list(df.columns) == STANDARD_COLUMNS

    def test_reads_all_four_tabs_and_ignores_dashboard_cells(self, df):
        assert len(df) == TOTAL_ROWS

    def test_activity_date_is_a_date_and_rows_are_sorted(self, df):
        assert all(type(d) is date for d in df["activity_date"])
        assert list(df["activity_date"]) == sorted(df["activity_date"])

    def test_source_row_is_renumbered_from_one(self, df):
        assert list(df["source_row"]) == list(range(1, TOTAL_ROWS + 1))

    def test_input_forms_are_equivalent(self, xlsx_path, xlsx_bytes):
        from_path = load_transactions_xlsx(xlsx_path)
        from_str = load_transactions_xlsx(str(xlsx_path))
        from_bytes = load_transactions_xlsx(xlsx_bytes)
        from_buffer = load_transactions_xlsx(io.BytesIO(xlsx_bytes))
        for other in (from_str, from_bytes, from_buffer):
            pd.testing.assert_frame_equal(from_path, other)

    def test_non_ledger_sheet_is_skipped(self, tmp_path):
        wb = _build_workbook()
        notes = wb.create_sheet("Notes")
        notes.append(["just", "some", "notes"])
        path = tmp_path / "with_notes.xlsx"
        wb.save(path)
        assert len(load_transactions_xlsx(path)) == TOTAL_ROWS

    def test_rows_after_the_first_blank_date_are_ignored(self, tmp_path):
        wb = _build_workbook()
        ws = wb["Other Options"]
        ws.append([None] * 11)  # blank spacer row
        ws.append(list(OTHER_OPTION_ROWS[0]))  # stray row after the block
        path = tmp_path / "trailing.xlsx"
        wb.save(path)
        assert len(load_transactions_xlsx(path)) == TOTAL_ROWS

    def test_workbook_without_a_ledger_sheet_raises(self, tmp_path):
        wb = Workbook()
        wb.active.append(["not", "a", "ledger"])
        path = tmp_path / "empty.xlsx"
        wb.save(path)
        with pytest.raises(ValueError, match="trade"):
            load_transactions_xlsx(path)

    def test_non_xlsx_bytes_raise_value_error(self):
        with pytest.raises(ValueError):
            load_transactions_xlsx(b"this is plainly not a workbook")


# ---------------------------------------------------------------------------
# Loader: row semantics
# ---------------------------------------------------------------------------

class TestLoaderSemantics:
    def test_transaction_types_collapse_to_buy_sell_and_pass_through_others(self, df):
        assert set(df["transaction_type"]) == {
            "Buy", "Sell", "Cash Settlement", "Dividend", "Other Fee",
        }

    def test_open_and_close_variants_map_to_buy_and_sell(self, df):
        assert _row(df, symbol=ETRADE_SHORT)["transaction_type"] == "Sell"  # Sell To Open
        assert _row(df, symbol=ETRADE_CLOSE)["transaction_type"] == "Buy"  # Buy To Close

    def test_quantity_is_signed_buy_positive_sell_negative(self, df):
        assert _row(df, symbol=IBKR_SHORT)["quantity"] == -1  # Sell
        assert _row(df, symbol=IBKR_LONG)["quantity"] == 1  # Buy
        assert _row(df, symbol=ETRADE_SHORT)["quantity"] == -2  # Sell To Open
        assert _row(df, symbol="BETA")["quantity"] == -20  # stock sell

    def test_quantity_is_empty_for_non_trade_rows(self, df):
        non_trade = df[df["transaction_type"].isin(["Cash Settlement", "Dividend", "Other Fee"])]
        assert non_trade["quantity"].isna().all()

    def test_net_amount_is_taken_verbatim(self, df):
        assert _row(df, symbol=IBKR_SHORT)["net_amount"] == pytest.approx(49.00)
        assert _row(df, symbol=IBKR_LONG)["net_amount"] == pytest.approx(-21.00)
        assert _row(df, symbol=ASSIGNED)["net_amount"] == pytest.approx(-25.00)

    def test_commission_is_a_magnitude(self, df):
        assert _row(df, symbol=OTHER_OPTION)["commission"] == pytest.approx(1.50)
        assert _row(df, symbol="BETA")["commission"] == pytest.approx(5.00)

    def test_gross_adds_the_commission_back_to_net(self, df):
        # 1 contract x 0.50 x 100 = 50.00 credit; 1 x 0.20 x 100 = 20.00 debit.
        assert _row(df, symbol=IBKR_SHORT)["gross_amount"] == pytest.approx(50.00)
        assert _row(df, symbol=IBKR_LONG)["gross_amount"] == pytest.approx(-20.00)
        assert _row(df, symbol="BETA")["gross_amount"] == pytest.approx(1000.00)

    def test_non_trade_rows_have_gross_equal_net_and_zero_commission(self, df):
        for typ in ("Dividend", "Other Fee", "Cash Settlement"):
            row = df[df["transaction_type"] == typ].iloc[0]
            assert row["gross_amount"] == pytest.approx(row["net_amount"])
            assert row["commission"] == 0.0


class TestLoaderSymbolsAndAccounts:
    def test_ibkr_occ_symbols_pass_through(self, df):
        assert IBKR_SHORT in set(df["symbol"])
        assert OTHER_OPTION in set(df["symbol"])

    def test_etrade_option_symbols_are_rebuilt_as_occ(self, df):
        symbols = set(df["symbol"])
        assert ETRADE_SHORT in symbols
        assert ETRADE_LONG in symbols
        assert not any(s.startswith("SPXW MAR") for s in symbols)

    def test_stock_and_fee_symbols_are_kept(self, df):
        symbols = set(df["symbol"])
        assert {"ACME", "BETA", "FUND", "-"} <= symbols

    def test_etrade_rows_use_the_default_virtual_account(self, df):
        etrade = df[df["symbol"].isin([ETRADE_SHORT, "ACME", "BETA"])]
        assert set(etrade["account_id"]) == {"E*Trade"}

    def test_etrade_rows_use_a_supplied_account_id(self, xlsx_path):
        out = load_transactions_xlsx(xlsx_path, account_id="000-000000-000")
        assert set(out[out["symbol"] == "ACME"]["account_id"]) == {"000-000000-000"}

    def test_ibkr_rows_keep_their_masked_account(self, df):
        assert set(df[df["symbol"] == OTHER_OPTION]["account_id"]) == {"U***12345"}

    def test_description_is_kept(self, df):
        assert _row(df, symbol=IBKR_SHORT)["description"] == "SPXW 03MAR25 5000 P"


# ---------------------------------------------------------------------------
# MCP routing
# ---------------------------------------------------------------------------

class TestMcpRouting:
    def test_load_and_merge_by_path(self, xlsx_path):
        merged, balances, warnings = _load_and_merge(paths=[str(xlsx_path)])
        assert len(merged) == TOTAL_ROWS
        assert balances == []
        assert warnings == []
        assert set(merged["account_id"]) == {"U***12345", "E*Trade"}

    def test_load_and_merge_by_base64_content(self, xlsx_bytes):
        fc = [{"name": "ledger.xlsx", "data_base64": base64.b64encode(xlsx_bytes).decode()}]
        merged, _balances, warnings = _load_and_merge(file_contents=fc)
        assert len(merged) == TOTAL_ROWS
        assert warnings == []

    def test_xlsm_extension_is_routed_too(self, tmp_path, xlsx_bytes):
        path = tmp_path / "ledger.xlsm"
        path.write_bytes(xlsx_bytes)
        merged, _balances, warnings = _load_and_merge(paths=[str(path)])
        assert len(merged) == TOTAL_ROWS
        assert warnings == []

    def test_csv_account_id_aligns_etrade_rows(self, xlsx_path):
        df, _bal, warn = _load_file_from_path(str(xlsx_path), csv_account_id="000-000000-000")
        assert warn is None
        assert set(df[df["symbol"] == "ACME"]["account_id"]) == {"000-000000-000"}

    def test_text_content_is_a_warning_not_a_crash(self):
        fc = [{"name": "ledger.xlsx", "data_text": "plain text, not a workbook"}]
        merged, _balances, warnings = _load_and_merge(file_contents=fc)
        assert merged.empty
        assert len(warnings) == 1 and warnings[0].startswith("ledger.xlsx")

    def test_missing_file_is_a_warning(self, tmp_path):
        merged, _balances, warnings = _load_and_merge(paths=[str(tmp_path / "gone.xlsx")])
        assert merged.empty
        assert len(warnings) == 1 and warnings[0].startswith("gone.xlsx")

    def test_loading_with_a_qfx_warns_about_double_counting(self, xlsx_path, tmp_path):
        qfx = tmp_path / "statement.qfx"
        qfx.write_text(SYNTHETIC_QFX, encoding="latin-1")
        _merged, _balances, warnings = _load_and_merge(paths=[str(xlsx_path), str(qfx)])
        assert any("xlsx" in w.lower() and "double" in w.lower() for w in warnings)

    def test_xlsx_alone_does_not_warn_about_double_counting(self, xlsx_path):
        _merged, _balances, warnings = _load_and_merge(paths=[str(xlsx_path)])
        assert not any("double" in w.lower() for w in warnings)


# ---------------------------------------------------------------------------
# MCP tools end to end
# ---------------------------------------------------------------------------

class TestMcpTools:
    def test_get_transaction_summary_reads_the_full_ledger(self, xlsx_path):
        result = get_transaction_summary(paths=[str(xlsx_path)])
        assert "error" not in result
        assert result["total_rows"] == TOTAL_ROWS
        assert result["date_range"] == {"start": "2025-03-03", "end": "2025-03-24"}
        assert set(result["accounts"]) == {"U***12345", "E*Trade"}
        assert result["transaction_types"]["Cash Settlement"] == 1
        assert result["transaction_types"]["Dividend"] == 1

    def test_compute_daily_pnl_keeps_only_spx_spxw_option_rows(self, xlsx_path):
        result = compute_daily_pnl(paths=[str(xlsx_path)], initial_capital=10000.0)
        assert "error" not in result
        # 5 SPXW option trades + the SPXW assignment settlement; the other
        # option, the stocks, the dividend and the fee are outside the strategy.
        assert result["total_rows"] == 6
        expected = 49.00 - 21.00 + 78.00 - 21.00 - 30.50 - 25.00
        assert result["total_realized_pnl"] == pytest.approx(expected)

    def test_expire_inference_depends_on_signed_quantity(self, xlsx_path):
        # Both short legs (IBKR 5000P on 3/3, E*Trade 5100P on 3/10) are sold
        # on their expiry day with no same-contract buyback, so the engine
        # infers expiry for each.  That only fires when Sell quantity < 0.
        result = compute_daily_pnl(paths=[str(xlsx_path)], initial_capital=10000.0)
        assert result["expire_inferred_count"] == 2


# ---------------------------------------------------------------------------
# E*Trade rows versus the E*Trade CSV loader
# ---------------------------------------------------------------------------

# The same five trades in both shapes.  The CSV carries unsigned magnitudes
# (the E*Trade loader signs them); the workbook records positive quantity,
# signed net and negative commission.
SYNTHETIC_ETRADE_CSV = (
    "Trade Date,Order Type,Security,Cusip,Transaction Description,Quantity,"
    "Executed Price,Commission,Net Amount\n"
    "3/10/2025,Sell To Open,SPXW MAR 10 '25 $5100 PUT,,"
    "PUT  SPXW   03/10/25  5100.000,2,0.40,1.00,78.00\n"
    "3/10/2025,Buy Open,SPXW MAR 10 '25 $5050 PUT,,"
    "PUT  SPXW   03/10/25  5050.000,2,0.10,1.00,21.00\n"
    "3/12/2025,Buy Open,SPXW MAR 12 '25 $5200 CALL,,"
    "CALL  SPXW   03/12/25  5200.000,1,0.20,0.50,20.50\n"
    "3/5/2025,Buy,ACME,000000000,ACME CORP UNSOLICITED TRADE,2,100,0.0000,200\n"
    "3/18/2025,Sell,BETA,000000001,BETA INC UNSOLICITED TRADE,20,50,5.00,995.00\n"
)
ETRADE_AS_WORKBOOK = [
    (date(2025, 3, 10), *ETRADE, "SPXW", "SPXW MAR 10 '25 $5100 PUT",
     "PUT  SPXW   03/10/25  5100.000", "Sell To Open", 2, 0.40, -1.00, 78.00),
    (date(2025, 3, 10), *ETRADE, "SPXW", "SPXW MAR 10 '25 $5050 PUT",
     "PUT  SPXW   03/10/25  5050.000", "Buy Open", 2, 0.10, -1.00, -21.00),
    (date(2025, 3, 12), *ETRADE, "SPXW", "SPXW MAR 12 '25 $5200 CALL",
     "CALL  SPXW   03/12/25  5200.000", "Buy Open", 1, 0.20, -0.50, -20.50),
    (date(2025, 3, 5), *ETRADE, "ACME", "ACME",
     "ACME CORP UNSOLICITED TRADE", "Buy", 2, 100, 0, -200.00),
    (date(2025, 3, 18), *ETRADE, "BETA", "BETA",
     "BETA INC UNSOLICITED TRADE", "Sell", 20, 50, -5.00, 995.00),
]
IDENTITY_COLUMNS = [
    "activity_date", "symbol", "transaction_type", "quantity", "price",
    "net_amount",
]


@pytest.fixture
def etrade_frame(tmp_path) -> pd.DataFrame:
    wb = Workbook()
    _sheet(wb, "Index Options", ETRADE_AS_WORKBOOK, first=True)
    path = tmp_path / "etrade_only.xlsx"
    wb.save(path)
    return load_transactions_xlsx(path)


class TestEtradeRowsVersusEtradeCsvLoader:
    def test_same_trades_give_identical_identity_columns(self, etrade_frame):
        from_csv = load_transactions_etrade_csv(SYNTHETIC_ETRADE_CSV.encode("utf-8"))

        def ordered(frame):
            return (
                frame[IDENTITY_COLUMNS]
                .sort_values(["activity_date", "symbol"])
                .reset_index(drop=True)
            )

        pd.testing.assert_frame_equal(
            ordered(etrade_frame), ordered(from_csv), check_dtype=False
        )

    def test_commission_is_the_workbook_column_as_recorded(self, etrade_frame):
        # 2 contracts x 0.40 x 100 = 80.00 but 78.00 was received.  The E*Trade
        # CSV loader reports that 2.00 gap as commission; the workbook records
        # E*Trade's own charge (exchange fees are already in the recorded
        # figures), so its column is used as is and the two paths
        # intentionally differ.
        sto = etrade_frame[etrade_frame["symbol"] == ETRADE_SHORT].iloc[0]
        assert sto["commission"] == pytest.approx(1.00)
        assert sto["gross_amount"] == pytest.approx(78.00 + 1.00)


class TestOptionalDependency:
    def test_server_still_imports_and_the_loader_reports_a_clear_error_without_openpyxl(self):
        script = """
import sys
sys.modules["openpyxl"] = None  # makes `import openpyxl` raise ImportError
import spx_trade_desk.mcp.server
from spx_trade_desk.tradelog.io.load_xlsx import load_transactions_xlsx
try:
    load_transactions_xlsx(b"not a workbook")
except ValueError as exc:
    print("CLEAR_ERROR", "openpyxl" in str(exc))
"""
        result = subprocess.run(
            [sys.executable, "-c", script],
            cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True)
        assert result.returncode == 0, result.stderr
        assert "CLEAR_ERROR True" in result.stdout


class TestLedgerRobustness:
    """Review fixes: header shadowing and silently dropped rows."""

    def test_a_dashboard_header_with_a_ledger_name_does_not_shadow_the_ledger_column(self, tmp_path):
        wb = Workbook()
        ws = _sheet(wb, "Index Options", INDEX_ROWS, first=True)
        ws["M1"] = "Account"          # a dashboard label that repeats a ledger header
        ws["M2"] = "dashboard value"
        path = tmp_path / "shadow.xlsx"
        wb.save(path)
        out = load_transactions_xlsx(path)
        assert "" not in set(out["account_id"])  # the dashboard cell would have read as a blank account
        assert set(out["account_id"]) == {IBKR[1], "E*Trade"}

    def test_rows_with_an_invalid_date_are_counted_logged_and_flagged(self, tmp_path, caplog):
        bad = [
            (46000, *IBKR, "SPXW", IBKR_SHORT, "SPXW 03MAR25 5000 P", "Sell", 1, 0.50, -1.00, 49.00),
            ("2025-13-45", *IBKR, "SPXW", IBKR_SHORT, "SPXW 03MAR25 5000 P", "Sell", 1, 0.50, -1.00, 49.00),
        ]
        wb = Workbook()
        _sheet(wb, "Index Options", INDEX_ROWS + bad, first=True)
        path = tmp_path / "bad_dates.xlsx"
        wb.save(path)
        with caplog.at_level("WARNING"):
            out = load_transactions_xlsx(path)
        assert len(out) == len(INDEX_ROWS)
        assert out.attrs["skipped_rows"] == 2
        assert "skipped 2 row" in out.attrs["warning"]
        assert any("skipped 2 row" in r.getMessage() for r in caplog.records)

    def test_a_clean_workbook_carries_no_warning(self, xlsx_path):
        out = load_transactions_xlsx(xlsx_path)
        assert "warning" not in out.attrs and "skipped_rows" not in out.attrs

    def test_the_mcp_loader_returns_the_skip_warning(self, tmp_path):
        from spx_trade_desk.mcp.server import _load_file_from_path
        bad = [(46000, *IBKR, "SPXW", IBKR_SHORT, "SPXW 03MAR25 5000 P", "Sell", 1, 0.50, -1.00, 49.00)]
        wb = Workbook()
        _sheet(wb, "Index Options", INDEX_ROWS + bad, first=True)
        path = tmp_path / "bad_dates.xlsx"
        wb.save(path)
        df, balance, warning = _load_file_from_path(str(path))
        assert balance is None and len(df) == len(INDEX_ROWS)
        assert warning and "skipped 1 row" in warning
