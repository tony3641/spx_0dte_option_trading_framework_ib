"""
Consolidated trade-log workbook (``.xlsx``) parser.

Parses the multi-tab ledger that merges IBKR and E*Trade trades into one
workbook (tabs ``Index Options`` / ``Other Options`` / ``Stock & ETF`` /
``Other Transactions``) and returns a DataFrame with the same column schema
as the other loaders:

    activity_date, account_id, description, transaction_type, symbol,
    quantity, price, gross_amount, commission, net_amount, source_row

Every tab shares one header::

    Date, Source, Account, Underlying, Symbol, Description, Order Type,
    Quantity, Price, Commission, Net Amount

Columns are looked up by header name, so dashboard cells that sit beside the
ledger (further columns on the first tab) are ignored, and a tab without that
header is skipped.  A tab's data block ends at the first blank ``Date``.

Conventions of the sheet versus the standard frame:

* **Quantity is always positive** in the sheet; the standard frame carries it
  signed (``Buy*`` +, ``Sell*`` -), which the P&L engine's expire-inference
  relies on.  Rows with no quantity (cash settlement, dividend, fee) stay
  empty.
* **Net Amount is already signed** and is taken verbatim, never recomputed.
* **Commission is negative** in the sheet; the frame carries its magnitude as
  recorded and ``gross_amount`` is net with that cost added back.  IBKR's
  column includes all fees; E*Trade's is only the broker's charge, because
  exchange fees are already in the recorded fill price.  (The E*Trade CSV
  loader instead reports the whole gap between ``quantity x price`` and the
  net as commission, so commission totals differ between the two paths.)
* **Order Type** ``Buy*`` / ``Sell*`` (Buy Open, Sell To Open, Buy To Close,
  ...) collapses to ``Buy`` / ``Sell``; anything else (Cash Settlement,
  Dividend, Other Fee) passes through under its own name, which is what the
  P&L engine keys on.
* **E*Trade option symbols are not OCC** (``SPXW MAR 10 '25 $5100 PUT``); they
  are rebuilt from ``Description`` with the same pattern the E*Trade CSV
  loader uses.  IBKR symbols are already OCC and pass through.
* **Account** is the sheet's value, except E*Trade rows, which get the
  ``account_id`` argument (default: the virtual ``"E*Trade"`` account) so they
  line up with the E*Trade CSV / PDF loaders.  The sheet's IBKR account is
  masked (``U***12345``) and is kept as is.
"""

from __future__ import annotations

import io
import logging
import math
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO, Optional, Union

import pandas as pd

from spx_trade_desk.tradelog.domain.parse_option_symbol import build_occ_symbol
from spx_trade_desk.tradelog.io.load_etrade_csv import (
    ETRADE_ACCOUNT_ID,
    _OPT_PATTERN,
    _parse_mdyy,
)

log = logging.getLogger(__name__)

LEDGER_HEADER = [
    "Date", "Source", "Account", "Underlying", "Symbol", "Description",
    "Order Type", "Quantity", "Price", "Commission", "Net Amount",
]

OUTPUT_COLUMNS = [
    "activity_date", "account_id", "description", "transaction_type",
    "symbol", "quantity", "price", "gross_amount", "commission",
    "net_amount", "source_row",
]

# The header sits in row 1 of the real file; allow a little slack above it.
_MAX_HEADER_SCAN_ROWS = 30

_ETRADE_SOURCE = "etrade"


def _norm(cell) -> str:
    return str(cell).strip().lower() if cell is not None else ""


def _num(value) -> float:
    """Cell -> float; blank or non-numeric cells become NaN."""
    if value is None or value == "":
        return math.nan
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", ""))
    except ValueError:
        return math.nan


def _to_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        parsed = pd.to_datetime(value, errors="coerce")
        return None if pd.isna(parsed) else parsed.date()
    return None


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _open_workbook(file_or_path: Union[BinaryIO, bytes, str, Path]):
    """Open read-only with cached values; map bad input to ``ValueError``."""
    try:
        # Imported here, not at module level: openpyxl is only needed to read a workbook, and a missing
        # package must not stop the MCP server (which imports this module) from starting.
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ValueError('Reading an .xlsx workbook needs the openpyxl package: pip install "openpyxl>=3.1"') from exc
    if isinstance(file_or_path, (str, Path)):
        source = str(file_or_path)
    else:
        raw = file_or_path if isinstance(file_or_path, bytes) else file_or_path.read()
        source = io.BytesIO(raw)
    try:
        return load_workbook(source, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError(f"Not a readable .xlsx workbook: {exc}") from exc


def _find_header(rows: list[tuple]) -> Optional[tuple[int, dict[str, int]]]:
    """Return (header row index, {normalized column name: column index})."""
    required = {c.lower() for c in LEDGER_HEADER}
    for idx, row in enumerate(rows):
        cols = {_norm(cell): i for i, cell in enumerate(row) if _norm(cell)}
        if required.issubset(cols):
            return idx, cols
    return None


def _convert_row(get, account_id: str) -> Optional[dict]:
    """Convert one ledger row (``get(name)`` -> cell) to a standard-frame row."""
    activity_date = _to_date(get("Date"))
    if activity_date is None:
        return None

    source = _norm(get("Source"))
    is_etrade = source == _ETRADE_SOURCE
    order_type = str(get("Order Type") or "").strip()
    description = str(get("Description") or "").strip()
    symbol = str(get("Symbol") or "").strip()

    lowered = order_type.lower()
    if lowered.startswith("buy"):
        tx_type, direction = "Buy", +1
    elif lowered.startswith("sell"):
        tx_type, direction = "Sell", -1
    else:
        tx_type, direction = order_type or "Unknown", None

    quantity = _num(get("Quantity"))
    if direction is not None and not math.isnan(quantity):
        quantity = direction * abs(quantity)

    if is_etrade:
        match = _OPT_PATTERN.match(description)
        if match:
            expiry = _parse_mdyy(f"{match.group(3)}/{match.group(4)}/{match.group(5)}")
            if expiry is not None:
                symbol = build_occ_symbol(
                    match.group(2).strip(), expiry, match.group(1), float(match.group(6))
                )

    net = _num(get("Net Amount"))
    commission = abs(_num(get("Commission")))
    commission = 0.0 if math.isnan(commission) else commission
    gross = net + commission

    return {
        "activity_date": activity_date,
        "account_id": account_id if is_etrade else str(get("Account") or "").strip(),
        "description": description,
        "transaction_type": tx_type,
        "symbol": symbol,
        "quantity": quantity,
        "price": _num(get("Price")),
        "gross_amount": gross,
        "commission": commission,
        "net_amount": net,
    }


def load_transactions_xlsx(
    file_or_path: Union[BinaryIO, bytes, str, Path],
    account_id: str = ETRADE_ACCOUNT_ID,
) -> pd.DataFrame:
    """
    Parse a consolidated trade-log workbook.

    Parameters
    ----------
    file_or_path:
        A path (str / Path), raw bytes, or a binary file-like object.
    account_id:
        Account id to tag E*Trade rows with.  Defaults to the virtual
        ``"E*Trade"`` account; pass the real E*Trade statement account id
        when a PDF is loaded alongside so the merge layer dedups overlapping
        option trades.

    Returns
    -------
    pd.DataFrame
        Standard 11-column transaction schema, sorted by date, with every
        ledger tab's rows (options, stocks, dividends, fees) combined.

    Raises
    ------
    ValueError
        If the input is not a readable workbook, or no tab carries the
        ledger header.
    """
    wb = _open_workbook(file_or_path)
    rows: list[dict] = []
    found_ledger = False
    try:
        for ws in wb.worksheets:
            sheet_rows = list(ws.iter_rows(values_only=True))
            header = _find_header(sheet_rows[:_MAX_HEADER_SCAN_ROWS])
            if header is None:
                log.info("xlsx: skipping tab %r (no ledger header)", ws.title)
                continue
            found_ledger = True
            header_idx, cols = header
            for raw in sheet_rows[header_idx + 1:]:
                def get(name, raw=raw):
                    i = cols[name.lower()]
                    return raw[i] if i < len(raw) else None

                if _is_blank(get("Date")):
                    break  # end of this tab's data block
                row = _convert_row(get, account_id)
                if row is not None:
                    rows.append(row)
    finally:
        wb.close()

    if not found_ledger:
        raise ValueError(
            "Could not find a trade-log sheet in the workbook "
            f"(expected header: {', '.join(LEDGER_HEADER)})."
        )

    out = pd.DataFrame(rows, columns=OUTPUT_COLUMNS[:-1])
    out = out.sort_values("activity_date", kind="stable").reset_index(drop=True)
    out["source_row"] = range(1, len(out) + 1)
    return out
